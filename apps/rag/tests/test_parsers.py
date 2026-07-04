from io import BytesIO

import pytest
from docx import Document

from app.parsers import UnsupportedDocumentType, parse_document


def test_parse_markdown_as_plain_text():
    result = parse_document(
        filename="notes.md",
        content_type="text/markdown",
        data=b"# Title\n\nKnowledge base content.",
    )

    assert result.text == "# Title\n\nKnowledge base content."


def test_parse_pdf_named_text_plain_as_ocr_text():
    result = parse_document(
        filename="scanned.pdf",
        content_type="text/plain",
        data="OCR 提取后的文本".encode("utf-8"),
    )

    assert result.text == "OCR 提取后的文本"


def test_parse_docx_extracts_paragraph_text():
    doc = Document()
    doc.add_paragraph("First paragraph")
    doc.add_paragraph("Second paragraph")
    buffer = BytesIO()
    doc.save(buffer)

    result = parse_document(
        filename="brief.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        data=buffer.getvalue(),
    )

    assert result.text == "First paragraph\nSecond paragraph"


def test_parse_unsupported_doc_extension_fails_clearly():
    with pytest.raises(UnsupportedDocumentType) as exc:
        parse_document(
            filename="legacy.doc",
            content_type="application/msword",
            data=b"not-supported",
        )

    assert "legacy.doc" in str(exc.value)
