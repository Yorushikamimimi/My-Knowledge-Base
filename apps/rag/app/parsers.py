from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


class UnsupportedDocumentType(Exception):
    pass


class DocumentParseError(Exception):
    pass


@dataclass(frozen=True)
class ParsedDocument:
    text: str


TEXT_EXTENSIONS = {".txt", ".md", ".markdown"}


def parse_document(filename: str, content_type: str | None, data: bytes) -> ParsedDocument:
    normalized_content_type = (content_type or "").lower()
    if normalized_content_type.startswith("text/") or "markdown" in normalized_content_type:
        return ParsedDocument(text=data.decode("utf-8", errors="replace"))
    extension = Path(filename).suffix.lower()
    if extension in TEXT_EXTENSIONS:
        return ParsedDocument(text=data.decode("utf-8", errors="replace"))
    if extension == ".pdf":
        return ParsedDocument(text=_parse_pdf(data))
    if extension == ".docx":
        return ParsedDocument(text=_parse_docx(data))
    raise UnsupportedDocumentType(f"Unsupported document type for {filename}")


def _parse_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n".join(page.strip() for page in pages if page.strip())
    except Exception as exc:
        raise DocumentParseError("Failed to parse PDF document") from exc


def _parse_docx(data: bytes) -> str:
    try:
        document = Document(BytesIO(data))
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
        return "\n".join(paragraph for paragraph in paragraphs if paragraph)
    except Exception as exc:
        raise DocumentParseError("Failed to parse DOCX document") from exc
